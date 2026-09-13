from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('C:/Users/李/Desktop/26B主论文_标号统一版kk.docx')
OUT=ROOT/'outputs/论文修订/26B主论文_Q4插图与覆盖条件修订版.docx'
d=Document(SOURCE)
def find(prefix):
    return next(p for p in d.paragraphs if p.text.startswith(prefix))
def after(p,text='',template=None):
    e=OxmlElement('w:p'); p._p.addnext(e)
    n=Paragraph(e,p._parent)
    source=(template or p)._p.pPr
    if source is not None: e.append(deepcopy(source))
    if text: n.add_run(text)
    return n
def replace(p,text):
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if prop is not None:r._r.insert(0,prop)

p=find('问题四含全向与定向干扰源')
after(p,'根据附件1第2.2节，机器狗的检测点允许超出目标区域。因此，半径1800 m的圆域仅限定干扰源的分布范围，不限制机器狗的检测位置；B、C路线中的圆外测站符合题设。')
p=find('圆形距离覆盖不能排除')
replace(p,'圆形距离覆盖不能排除测站全部落在定向源背面的情况。为此，本问构造 B、C 两组认证测站及访问路线，并同时检验接收距离与方位覆盖。对任意源位置 X，先筛选与 X 距离不超过1000 m的测站，构成可达测站集合；1000 m为题设有效接收半径下界。记该集合的测站数为 m(X)，仅将 X 指向这些可达测站的方向角升序排列，并补充首角加360°形成循环序列。式（4-5）中的方向角均针对这一集合定义；当集合为空时，令最大循环角隙为360°。源位置与测站重合时可直接近距离检测并清除，单独处理。')
p=find('若所有测站方向均能')
replace(p,'若全部可达测站方向均能被某个开半圆容纳，则存在一种发射朝向使定向源对这些测站均背向；反之，若可达测站的最大循环角隙不超过180°，则任意前向半平面内至少包含一个距离不超过1000 m的测站，从而同时满足方向与接收距离条件。因此，认证路线应满足')
for t in d.tables:
    if '(4-5)' in ''.join(t._tbl.itertext()):
        for node in t._tbl.iter(qn('m:t')):
            if node.text=='1≤j≤m':node.text='1≤j≤m(X)'
p=find('程序在作业圆域内进行网格搜索')
for run in p.runs:
    if '网格搜索与边界加密。' in run.text:
        run.text=run.text.replace('网格搜索与边界加密。','网格搜索与边界加密，各源位置仅使用1000 m内的可达测站计算角隙。')
p=find('在作业圆域内对源位进行网格搜索')
for run in p.runs:
    run.text=run.text.replace('全部认证测站','距离不超过1000 m的可达认证测站')

caption_template=find('图 5-8 J+')
def figure(anchor,name,caption,intro):
    ref=after(anchor,intro)
    p=after(ref)
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(5)
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT/f'outputs/q4/figures/paper_geometry/{name}.png'),width=Cm(15))
    c=after(p,caption,caption_template)
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent=Cm(0)
    c.paragraph_format.keep_with_next=False
    c.paragraph_format.space_after=Pt(6)
    for r in c.runs:r.font.size=Pt(10.5)
figure(find('由于保守可行域包含真实源'),'q4_safe_pruning',
       '图 5-9 基于保守可行域的共观测距离剪枝机制',
       '图5-9利用正式日志中的两次成功示向重建可行域，并以路线测站展示距离门：最短距离约950 m时保留候选共观测，约1974 m时安全跳过。候选站点用于解释判据，并非实际剪枝记录；保留也不意味着必然接收到信号。')
figure(find('C 路线比 B 多 2 个测站'),'q4_certified_routes',
       '图 5-10 B、C认证测站布局与开放访问路线',
       '图5-10给出了冻结的B、C测站布局及访问方向。虚线圆表示干扰源分布区域，星形与方形分别表示起点和终点；检测点可位于目标区域外，路线均为不返回起点的开放路径。')
OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
print(OUT)
