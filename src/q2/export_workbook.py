"""将问题2四类CSV结果导出为国赛论文附件用Excel工作簿。"""
from pathlib import Path
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT=Path(__file__).resolve().parents[2]; SRC=ROOT/'outputs'/'q2'/'tables'; OUT=SRC/'Q2_第二检测点实验.xlsx'
SHEETS=[('候选点评分','Q2候选点评分.csv','候选第二检测点的Q1接口联合评分'),('策略对比','Q2策略对比.csv','联合评分、最近点、固定90度点与随机点的离线回放对比'),('权重敏感性','Q2权重敏感性.csv','各权重±10%扰动后的最优点稳定性'),('圆域逼近检验','Q2圆域逼近检验.csv','内接正多边形近似半径1800m圆盘的误差上界')]
wb=Workbook(); wb.remove(wb.active)
for name,file,desc in SHEETS:
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    with open(SRC/file,encoding='utf-8-sig',newline='') as f: rows=list(csv.reader(f))
    ws['A1']='问题2 第二检测点实验'; ws['A1'].font=Font(name='Microsoft YaHei',size=16,bold=True,color='17365D')
    ws['A2']=desc; ws['A2'].font=Font(name='Microsoft YaHei',size=10,color='536579')
    for j,v in enumerate(rows[0],1):
        c=ws.cell(4,j,v); c.fill=PatternFill('solid',fgColor='17365D'); c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center')
    for r,row in enumerate(rows[1:],5):
        for c,v in enumerate(row,1): ws.cell(r,c,v).alignment=Alignment(horizontal='center')
    end=4+len(rows); ref=f'A4:{chr(64+len(rows[0]))}{end}'; tab=Table(displayName='T'+name.replace(' ','')[:20],ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(tab)
    for col in range(1,len(rows[0])+1): ws.column_dimensions[chr(64+col)].width=18
    ws.freeze_panes='A5'
wb.save(OUT); print(OUT)
