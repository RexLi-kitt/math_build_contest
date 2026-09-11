"""生成与 Q2方案卡一致的论文建模大纲 DOCX。"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent; ROOT=HERE.parents[1]; OUT=ROOT/'outputs'/'q2'/'问题2建模大纲_方案卡版.docx'

def set_cell(cell, fill=None, bold=False):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        for r in p.runs: r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.bold=bold
    if fill:
        tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers):
        c.text=v; set_cell(c,'17365D',True)
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for c,v in zip(cells,row): c.text=str(v); set_cell(c,'EAF2F8' if i%2==0 else None)
    doc.add_paragraph()
def h(doc,text,level=1): doc.add_heading(text,level=level)
def p(doc,text,boldlead=None):
    par=doc.add_paragraph(); par.paragraph_format.space_after=Pt(6)
    if boldlead:
        r=par.add_run(boldlead); r.bold=True; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    r=par.add_run(text); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); return par

doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Cm(2.2); sec.left_margin=sec.right_margin=Cm(2.4)
styles=doc.styles; styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('问题2 四指标主动第二检测点选择模型').font.color.rgb=RGBColor(0,0,0)
p(doc,'在首次检测返回 direction 的条件下，本模型为已知存在的一台全向干扰源选择第二检测点。模型以独立物理可行域、保守接收保障、定位精度与移动时间为主线，先粗搜定位潜在优区，再局部加密并从高分候选区域中选取执行点。')
h(doc,'1 建模边界与输入数据')
p(doc,'问题1的纯 AOA 半平面交不作为数值结果输入；问题2只复用角度误差带构造方法，并在本问独立加入源分布、首次 direction 距离和接收规则。')
table(doc,['数据','符号','作用'],[['首测点与示向度','S1, theta1','构造首测角度带与移动起点'],['示向度误差','epsilon=1 degree','两次测向的有界误差'],['源位置范围','||G||<=1800 m','限制干扰源，不限制机器狗'],['首次 direction','5<||G-S1||<=1500 m','构造 A1'],['接收半径','r in [1000,1500] m','1000 m 是保证接收半径'],['速度与清除阈值','v=5 m/s, MEC<=20 m','时间计算与清除判据']])
h(doc,'2 首测物理可行域 A1')
p(doc,'A1 = {G: |wrap(theta(G,S1)-theta1)|<=1 degree, ||G||<=1800, 5<||G-S1||<=1500}。首次 direction 不代表已经定位成功，也不能增加 ||G-S1||>=1000 的错误约束。')
p(doc,'在 A1 内作均匀面积近似采样，样本既用于几何与接收指标的期望估计，也用于离线回放；5m 内孔在样本判定中严格排除。')
h(doc,'3 两阶段候选点搜索')
p(doc,'粗网格 C 以 S1 为圆心枚举距离与方位，目的仅是识别全局高分中心。保留前两个中心后，在距离和方位邻域内生成局部加密候选集 C_ref。最终的归一化、评分、候选区域和执行点只能在 C_ref 内计算。')
p(doc,'候选点不与 ||S2||<=1800 相交；1800m 圆只约束源位置。规划距离上限 1800m 是计算搜索半径，不是机器狗禁区。')
h(doc,'4 四指标与归一化评分')
table(doc,['指标','计算口径','偏好'],[['几何质量 Qgeo','样本交会角 sin²(alpha) 的均值','越大越好'],['保证接收 Prec','||Gj-S2||<=1000m 的样本比例','越大越好'],['移动时间 Tmove','||S2-S1||/5','越小越好'],['预测半径 Rpred','仅保证接收样本的 A2 MEC 平均值','越小越好']])
p(doc,'对四项指标在 C_ref 上 min-max 归一化。综合评分为 J = w1 Qgeo_tilde + w2 Prec_tilde - w3 Tmove_tilde - w4 Rpred_tilde，其中 wi>=0 且四项之和为1。分母为0的候选点无保证接收样本，直接淘汰。')
h(doc,'5 高分候选区域与执行点')
p(doc,'令 J* 为 C_ref 中最大评分，取 eta=0.10。定义 R_good={s in C_ref: J(s)>=(1-eta)J*}。这保留对采样和观测误差更稳健的一组构型，而不是压缩为唯一峰值点。执行点为 R_good 内移动时间最短的点；时间并列时依次选预测 MEC 更小、评分更高的点。')
h(doc,'6 权重校准与在线闭环')
p(doc,'代码中的示例权重仅用于复现。论文最终权重采用外层蒙特卡洛与单纯形网格搜索：随机生成真源位置、接收半径和角度误差，以定位半径、接收失败率和总时间的加权损失选择最优稳定权重，并作各权重±10%扰动检验。')
p(doc,'二测返回 direction 时更新 A2=A1∩B2，MEC不超过20m才前往最小外接圆圆心附近 clear；返回 near 时直接 clear；返回 no_signal 时只排除以 S2 为圆心的1000m保证接收圆盘，不能将1000--1500m的无信号硬删除。')
h(doc,'7 实验设计与交付')
table(doc,['检验','输出'],[['几何与边界','A1、A2、近平行与清除阈值边界案例'],['接收逻辑','1000m内 no_signal 为零；1200m无信号不得硬删除'],['策略比较','四指标联合、仅几何、最近点、随机点'],['候选区域','R_good 数据表、峰值点和最终执行点'],['鲁棒性','权重±10%扰动后的执行点、接收率与预测MEC']])
fig=ROOT/'outputs'/'q2'/'figures'/'q2_four_metric_score.png'
if fig.exists():
    doc.add_picture(str(fig),width=Cm(13.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=doc.add_paragraph('图 1  局部加密候选集的四指标评分与高分候选区域'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
